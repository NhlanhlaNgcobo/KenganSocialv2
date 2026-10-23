from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'docs' / 'KenganSocial_Deployment_Checklist.docx'
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.7), Inches(8.3)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.05
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(3)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.styles['Caption'].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Caption'].font.bold = False
doc.styles['Caption'].font.size = Pt(8)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('KenganSocial deployment checklist  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(8)

COMPLETED = {1, 4, 7, 9, 10}
PARTIAL = {2, 3, 6, 8, 11, 14, 15, 16}

def check(p, ident):
    xml = f'''<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"><w:sdtPr><w:alias w:val="Task {ident} completed"/><w:id w:val="{1000+ident}"/><w14:checkbox><w14:checked w14:val="0"/><w14:checkedState w14:val="2612" w14:font="Segoe UI Symbol"/><w14:uncheckedState w14:val="2610" w14:font="Segoe UI Symbol"/></w14:checkbox></w:sdtPr><w:sdtContent><w:r><w:rPr><w:rFonts w:ascii="Segoe UI Symbol" w:hAnsi="Segoe UI Symbol"/></w:rPr><w:t>☐</w:t></w:r></w:sdtContent></w:sdt>'''
    if ident in COMPLETED:
        xml = xml.replace('w14:val="0"', 'w14:val="1"').replace(chr(0x2610), chr(0x2612))
    p._p.append(parse_xml(xml))

def item(n, title, action, done):
    p = doc.add_paragraph(style='Heading 2'); check(p,n); p.add_run(f'  {n:02d}  {title}')
    p = doc.add_paragraph(action); p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.add_run('Complete when  ').bold = True; p.add_run(done); p.paragraph_format.keep_with_next = True
    status = 'Verified locally' if n in COMPLETED else 'Partly implemented' if n in PARTIAL else 'Pending verification'
    p = doc.add_paragraph(f'{status}  |  Owner __________  Evidence _____________________')
    p.paragraph_format.space_after = Pt(7)
    for r in p.runs: r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string('555555')

def page(title, intro):
    doc.add_page_break(); doc.add_heading(title,1); doc.add_paragraph(intro)

doc.add_paragraph('KenganSocial deployment checklist', 'Title')
doc.add_paragraph('Release preparation  |  23 October 2026')
doc.add_paragraph('Current decision: public launch is blocked. Complete the required fixes and staging tests below before approving release. Five fixes are checked against local evidence. Partial items stay open. See the progress notes on page 5; staging acceptance remains required.')
doc.add_paragraph('Release owner ____________________  Target date ____________________')
doc.add_heading('1 Security fixes before any pilot',1)
item(1,'Restrict video inputs','Replace arbitrary video URLs and file paths with ownership-checked upload IDs. Restrict internal file reads and any permitted remote downloads.','Foreign uploads, paths outside the shared folder, private network URLs and unsafe redirects are rejected.')
item(2,'Enforce workspace roles','Add role checks to publishing, ad status, ad budgets, billing and other write routes. Membership alone must not authorize these actions.','Viewer writes return 403; permitted roles succeed; cross-workspace requests are denied.')
item(3,'Repair the export route','Validate the parent job and its workspace before creating an export. Implement the missing Python export endpoint or remove the obsolete API route.','Foreign and nonexistent jobs cannot be changed, and the supported export flow finishes successfully.')
item(4,'Reject missing workspaces','Require a valid workspace on every tenant route, including schedule drafts, edits, deletion and publishing.','Two users without workspaces cannot create or share records under a null workspace scope.')
item(5,'Protect private media','Separate private originals from public publishing assets. Use authenticated downloads or short-lived signed links for private files.','Anonymous and foreign-workspace downloads fail; public publishing links expire as designed.')

page('2 Billing and code quality','Required before charging customers. Baseline: the frontend build passed; lint had 4 warnings; backend type checking reported 23 errors. Audit findings need individual triage.')
item(6,'Implement paid access expiry','Choose recurring subscriptions or time-limited one-off purchases. Enforce expiry, renewal, failed payment and cancellation in API access checks.','Expired or cancelled access follows the agreed policy; successful renewal restores the correct allowance.')
item(7,'Make payment activation recoverable','Commit the payment marker and entitlement together, or add a durable recovery state and reconciliation.','Stopping the worker between receipt and activation, then retrying, grants access exactly once.')
item(8,'Make usage limits atomic','Replace separate quota reads and increments with an atomic reservation. Bound dependency waits and handle unavailable counters safely.','Concurrent requests cannot exceed the remaining cap; Redis failure cannot cause unmetered paid work.')
item(9,'Resolve dependency findings','Triage 47 server and 9 client audit findings, including 3 critical server findings. Upgrade safely and test affected features.','All critical and high production findings are fixed or have a documented, reviewed disposition.')
item(10,'Clear compiler errors and review warnings','Resolve all 23 backend diagnostics, including missing modules and upload types. Review 4 React Hook warnings and add frontend TS/TSX checks.','Backend and frontend type checks pass; lint has no errors and remaining warnings are reviewed.')
item(11,'Create repeatable release checks','Add CI for clean installs, builds, lint, type checks, selected editor tests and authorization/billing regression tests. Pin Python dependencies.','A clean checkout passes the checks using the same runtime and start command as production.')

page('3 Deployment and recovery','Required before opening staging to testers. Docker startup and the live cloud configuration were not verified during the local evaluation.')
item(12,'Validate the complete staging stack','Build the Linux API and video images; validate Caddy; start API, video service and Redis. Confirm the shared media volume works.','All services start successfully, HTTPS works, and Redis/video ports remain internal.')
item(13,'Configure production settings','Set production URLs, secrets, Firebase domains, provider callbacks and disabled development overrides. Deploy current Firebase rules and indexes.','The browser calls the intended API, CORS permits intended origins, and secrets are absent from source and images.')
item(14,'Verify storage and tenant queries','Test Firebase Storage read/write rules if used; separate read checks from upload-size checks. Filter calendar queries by workspace in Firestore.','Storage access passes role/tenant tests and calendar queries read only the requested workspace.')
item(15,'Persist uploads and recover jobs','Persist upload ownership/path metadata. Drain work on shutdown and test API, video and Redis restarts. Keep a single API instance until coordination is proven.','Uploads remain usable and interrupted jobs recover or fail clearly without duplicate publishing or charges.')
item(16,'Add readiness checks and alerts','Check Redis, Firestore and video-service readiness separately from liveness. Configure error, uptime, disk and resource monitoring.','Dependency failures change readiness and produce a test alert with an assigned responder.')
item(17,'Prove backups and rollback','Back up Firestore, media and required configuration. Restore into an isolated environment and rehearse deployment rollback.','Restored data and media are usable; recovery time, data loss window and rollback steps are recorded.')

page('4 Final verification and release','Use test users, test payments and approved provider test accounts. These checks establish release evidence; successful local unit tests do not replace them.')
item(18,'Test user and tenant journeys','Verify sign-up, consent, workspace creation, invitations, viewer restrictions and separate users in separate workspaces.','Each permitted journey succeeds and all negative authorization cases are denied.')
item(19,'Test media and publishing journeys','Upload a video, transcribe, edit and export it; schedule and publish through an approved provider test account.','The final media is playable and a scheduled post publishes once with the correct account and content.')
item(20,'Test payment lifecycle','Run initial payment, renewal, expiry, failed payment, cancellation, duplicate delivery and out-of-order webhooks in test mode.','Workspace access and billing records match every expected outcome, including restart recovery.')
item(21,'Validate the deployed site and capacity','Run scripts/check-deploy.mjs against staging domains. Check legal pages, HTTPS, CORS and browser API calls; exercise concurrent renders.','Deployment checks pass and response times, memory and disk remain within recorded operating limits.')
item(22,'Confirm provider and privacy readiness','Confirm required platform app approvals and permissions. Verify consent, export/deletion, retention, privacy disclosures and applicable POPIA arrangements.','Publicly advertised integrations are approved and the responsible owner records privacy and deletion verification.')
item(23,'Approve the release','Review completed evidence, refresh deployment instructions, assign the launch responder and record any deferred non-blocking improvements.','Required tasks pass; no unresolved security or paid-access blocker remains; the release owner signs below.')
doc.add_paragraph('Release decision   ☐ Approved   ☐ Hold     Release revision __________________')
doc.add_paragraph('Approved by __________________________  Date __________________________')
doc.add_paragraph('Deferred items and reasons __________________________________________________')
doc.add_paragraph('Reference: docs/DEPLOYMENT_EVALUATION_2026-10-23.md and its saved audit evidence. Findings describe revision d147ea7; verify the release revision again.',style='Caption')
page('5 Fixes and remaining release work','Implementation update 23 October 2026. Checked items reflect local regression evidence only. Public deployment remains on hold.')
for title, body in [
 ('Verified locally','01 Owned-upload inputs and Python file boundaries tested. 04 Missing workspace access denied. 07 Payment activation and its receipt marker commit atomically with retry coverage. 09 Both npm audits report zero vulnerabilities. 10 Both type checks pass and frontend lint is clean.'),
 ('Security and media progress','02 Viewer writes are denied; ad budgets, connections and billing require owner/admin. 03 The obsolete export endpoint returns 410; the complete editor journey still needs staging verification. 05 Private originals and anonymous media access remain unresolved.'),
 ('Recurring subscriptions','06 Checkout now uses recurring Paystack plans. Lifecycle events reconcile provider state; paid access expires at period end. Owners/admins can open Paystack subscription management. Plan changes and restarting a cancelled subscription currently require billing support. Existing one-off customers are not automatically enrolled.'),
 ('Reliability progress','08 Quota reservation is atomic; 30 concurrent attempts against four remaining uses permit exactly four. 14 Storage rules and workspace-scoped calendar queries updated; cloud/emulator verification pending. 15 Upload metadata survives registry recreation; worker draining and restart recovery remain open. 16 Dependency readiness implemented; operational alerts remain open.'),
 ('Build and test evidence','8 security/billing tests, 16 video-ad tests including real MP4 export, and 1 Python boundary test pass. Shared editor files match. Firebase Admin uses modular APIs; production tsx dependency and Node 24 Docker runtime updated. CI added but has not run on GitHub. Python dependency locking and Linux image validation remain open.'),
 ('Required before release','Complete private-media protection, staging and provider journeys, Paystack test-mode renewal/failure/cancellation, production settings, backups/restore, alerts, capacity checks and privacy/provider sign-off. No deployment, cloud rule change or live payment was performed.'),
]:
    doc.add_heading(title,2); doc.add_paragraph(body)
doc.add_paragraph('Detailed evidence and limitations: docs/DEPLOYMENT_FIX_PROGRESS_2026-10-23.md', style='Caption')

doc.core_properties.title='KenganSocial deployment checklist'
doc.core_properties.subject='Deployment fixes and release acceptance checklist'
doc.core_properties.author=''
doc.save(OUT)
print(OUT)
